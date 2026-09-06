# -*- coding: utf-8 -*-
"""
Bootstrap workbook για το portal-demo (κλώνος του Demo Portal).
Δημιουργεί ΟΛΑ τα απαιτούμενα sheet tabs με τη σωστή σειρά επικεφαλίδων,
όπως ακριβώς τις διαβάζει/γράφει ο κώδικας (src/*.js _HEADER_ORDER
constants) — όχι από ενδεχομένως ξεπερασμένο prose στο CLAUDE.md/SETUP.md.

Μετά τη δημιουργία: ανέβασμα σε Google Drive με αυτόματη μετατροπή σε
native Google Sheet (μέσω του συνδεδεμένου Google Drive MCP connector),
μετά μοίρασμα του Sheet με το email του service account (Editor) ώστε ο
Worker (portal-demo) να μπορεί να διαβάσει/γράψει.
"""
import openpyxl
from openpyxl.styles import Font, PatternFill

SHEETS = {
    "Άδειες": [
        "ID", "Timestamp", "EmployeeID", "EmployeeName", "TeamLeaderEmail",
        "Type", "StartDate", "EndDate", "Days", "Note", "Status",
        "DecisionDate", "DecisionNote", "EnteredBy",
        "FileKey", "FileName", "FileType",
        "GroupID",
    ],
    "Υπάλληλοι": [
        "EmployeeID", "Name", "Email", "TeamLeaderEmail", "AnnualDays", "Team", "HasDrivingLicense",
        "AMA", "AMKA", "Phone", "Address", "HireDate", "TechnicalTraining",
        "CvFileKey", "CvFileName", "CvFileType",
        "IdFileKey", "IdFileName", "IdFileType",
        "ResidencePermitFileKey", "ResidencePermitFileName", "ResidencePermitFileType",
        "CompanyPhone",
        "Status", "InactiveDate",
        "AFM",
        "HuskiesUsername", "HuskiesPassword",
        "PriorExperience",
        "FatherName", "IdNumber",
    ],
    "TeamLeaders": ["Name", "Email", "Team", "BackupEmail", "Away"],
    "Backoffice": ["Name", "Email"],
    "Directors": ["Email", "Name"],
    "Στόλος": [
        "ID", "Plate", "Type", "AssignedTo", "Status", "ServiceNote", "UpdatedBy", "UpdatedAt",
        "KteoDate", "ServiceDate", "Deductible", "CurrentMileage",
        "InsuranceCompany", "PolicyNumber", "InsuranceRenewalDate", "ServiceEnteredAt",
    ],
    "ΑτυχήματαΟχημάτων": [
        "ID", "VehicleID", "VehiclePlate", "EmployeeID", "EmployeeName", "Date", "Comments",
        "FileKey", "FileName", "FileType", "ReportedBy", "ReportedAt",
    ],
    "Εξοπλισμός": [
        "ID", "Name", "Category", "AssignedTo", "Status", "Note", "UpdatedBy", "UpdatedAt",
        "SerialNumber", "Model", "IMEI",
    ],
    "EPass": [
        "ID", "Label", "AssignedTechnicianId", "AssignedTechnicianName",
        "AssignedVehicleId", "AssignedVehiclePlate", "Status", "Note", "UpdatedBy", "UpdatedAt",
    ],
    "SIMs": ["ID", "CardNumber", "PIN", "PUK", "Tablet", "Note", "UpdatedBy", "UpdatedAt"],
    "Χρεώσεις": [
        "ID", "Type", "ItemID", "ItemLabel", "EmployeeID", "EmployeeName",
        "ChargedAt", "ReleasedAt", "ChargedBy", "ReleasedBy", "VehicleId", "VehiclePlate",
        "PickupMileage", "PickupPhotos", "ReturnMileage", "PickupDate", "ReturnDate",
        "OdometerPickupPhotoKey", "OdometerPickupPhotoName", "OdometerPickupPhotoType",
        "OdometerReturnPhotoKey", "OdometerReturnPhotoName", "OdometerReturnPhotoType",
    ],
    "Συνεργεία": ["Date", "CrewsJSON", "SavedBy", "SavedAt"],
    "Ενημερώσεις": [
        "ID", "Title", "Body", "FileKey", "FileName", "FileType",
        "PostedBy", "PostedByName", "PostedAt", "TargetTeams",
    ],
    "ΕρωτηματαΔιαθεσιμότητας": [
        "ID", "Question", "Deadline", "TargetTeams", "CreatedBy", "CreatedByName", "CreatedAt",
        "PublishAt", "EmailSentAt", "ResponseMode",
    ],
    "ΑπαντησειςΔιαθεσιμότητας": ["ID", "QueryID", "EmployeeID", "EmployeeName", "Answer", "AnsweredAt"],
    "ΕγγραφαΠαράδοσης": [
        "ID", "EmployeeID", "EmployeeName", "HoldingsJSON", "Status",
        "CreatedBy", "CreatedByName", "CreatedAt",
        "SignatureFileKey", "SignatureFileName", "SignatureFileType", "SignedAt",
    ],
    "Έγγραφα": [
        "ID", "EmployeeID", "EmployeeName", "DocType", "FieldsJSON",
        "CreatedBy", "CreatedByName", "CreatedAt",
    ],
    "AuditLog": ["Timestamp", "User", "Action", "Details"],
}

HEADER_FILL = PatternFill(start_color="1F3A5F", end_color="1F3A5F", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)

wb = openpyxl.Workbook()
wb.remove(wb.active)  # αφαίρεση του default "Sheet"

for name, headers in SHEETS.items():
    ws = wb.create_sheet(title=name)
    for col_idx, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col_idx, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
    ws.freeze_panes = "A2"
    # πλάτος στηλών ~ μήκος header (min 10, max 28)
    for col_idx, h in enumerate(headers, start=1):
        letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[letter].width = max(10, min(28, len(h) + 4))

out_path = "/sessions/dreamy-pensive-curie/mnt/outputs/portal-demo-sheet-bootstrap.xlsx"
wb.save(out_path)
print("Saved:", out_path)
print("Sheets:", len(SHEETS))
for n, h in SHEETS.items():
    print(f"  {n}: {len(h)} στήλες")
